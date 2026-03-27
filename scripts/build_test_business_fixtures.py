#!/usr/bin/env python3
"""Regenerate test_business_files as valid PDF/DOCX from UTF-8 text placeholders.

Run from repo root after installing deps (fpdf2, python-docx):
    python scripts/build_test_business_fixtures.py

Preserves full narrative under test_business_files/sources/*.txt; PDF stubs
reference those paths and carry enough Latin text for pdfplumber extraction
(>= parser minimum length).
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from fpdf import FPDF

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = REPO_ROOT / "test_business_files"
SOURCES_DIR = FIXTURE_DIR / "sources"


def make_stub_pdf(path: Path) -> None:
    """Write a minimal valid PDF with extractable Latin-only text (Helvetica)."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    lines = [
        "Graphsearch test fixture: valid PDF binary (not plain UTF-8 text).",
        "Full UTF-8 narrative is stored under test_business_files/sources/",
        "using the same base filename as this PDF with a .txt extension.",
        (
            "This paragraph exists so text extraction meets minimum length. "
            "Lorem ipsum dolor sit amet, consectetur adipiscing elit."
        ),
    ]
    w = pdf.epw
    for line in lines:
        pdf.multi_cell(w, 7, line)
    pdf.output(str(path))


def make_docx_from_text(path: Path, text: str) -> None:
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(str(path))


def main() -> int:
    if not FIXTURE_DIR.is_dir():
        print(f"Missing {FIXTURE_DIR}", file=sys.stderr)
        return 1

    SOURCES_DIR.mkdir(parents=True, exist_ok=True)

    for pdf_path in sorted(FIXTURE_DIR.glob("*.pdf")):
        raw = pdf_path.read_bytes()
        if raw.startswith(b"%PDF-"):
            print(f"skip (already PDF): {pdf_path.name}")
            continue
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError as e:
            print(f"skip (binary, not utf-8 text): {pdf_path.name} ({e})", file=sys.stderr)
            continue
        (SOURCES_DIR / f"{pdf_path.stem}.txt").write_text(text, encoding="utf-8")
        make_stub_pdf(pdf_path)
        print(f"wrote PDF + source: {pdf_path.name}")

    for docx_path in sorted(FIXTURE_DIR.glob("*.docx")):
        raw = docx_path.read_bytes()
        if raw.startswith(b"PK\x03\x04"):
            print(f"skip (already DOCX): {docx_path.name}")
            continue
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError as e:
            print(f"skip (binary): {docx_path.name} ({e})", file=sys.stderr)
            continue
        (SOURCES_DIR / f"{docx_path.stem}.txt").write_text(text, encoding="utf-8")
        make_docx_from_text(docx_path, text)
        print(f"wrote DOCX + source: {docx_path.name}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
